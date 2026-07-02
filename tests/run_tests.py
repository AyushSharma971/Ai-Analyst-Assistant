"""Local, offline, deterministic test suite — NO external calls.

Runs with the mock LLM, no OCR engine, no DB. Verifies the optimizations:
determinism, config-driven behavior (override a config -> behavior changes),
coercion, validation gates, risk rules from config, safe autofill formula
preservation, copilot grounding, and registry completeness.

    python tests/run_tests.py        # exits non-zero if any check fails
"""

from __future__ import annotations

import json
import os
import sys
import tempfile

# Make the project root importable when run as `python tests/run_tests.py`.
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.config import Settings
from app.contracts import IntakeMode, NevagInput
from app.document_ai import build_ocr_engine, parse_documents
from app.embeddings import MockEmbedding, build_embedding_provider
from app.extraction import coerce, extract_fields, load_extraction_prompt
from app.registry import AGENT_REGISTRY, describe_registry
from app.retrieval import chunk_documents, ingest_documents, retrieve
from app.service import NevagAgent
from app.validation import validate_fields
from app.vectorstore import (
    MemoryVectorStore,
    QdrantVectorStore,
    build_vector_store,
    validate_vector_store_config,
)
from scripts.make_fixtures import build_all

_OUT = os.path.join(tempfile.gettempdir(), "nevag_test_out")
_PATHS = build_all()
_RATER = _PATHS.pop("rater_template")
_FILES = [{"filename": os.path.basename(p), "path": p} for p in _PATHS.values()]

_RESULTS = []


def check(name, cond, detail=""):
    ok = bool(cond)
    _RESULTS.append(ok)
    print(("PASS" if ok else "FAIL"), "-", name, ("" if ok else f"  [{detail}]"))


def settings(**over):
    # Pin mock LLM + embeddings so tests are deterministic and never touch Ollama.
    base = {
        "llm_provider": "mock",
        "embedding_provider": "mock",
        "rater_template_path": _RATER,
        "rater_output_dir": _OUT,
    }
    base.update(over)  # callers may override
    return Settings(**base)


def run_pipeline(**over):
    return NevagAgent(settings=settings(**over)).run(
        NevagInput(query="New D&O submission", mode=IntakeMode.UPLOAD, files=list(_FILES))
    )


def write_json(name, obj):
    p = os.path.join(tempfile.gettempdir(), name)
    with open(p, "w", encoding="utf-8") as fh:
        json.dump(obj, fh)
    return p


# --------------------------------------------------------------------------- #
def test_determinism():
    a = run_pipeline()
    b = run_pipeline()
    check("determinism: same submission_id", a.submission_id == b.submission_id)
    check("determinism: same audit run_id", a.audit["run_id"] == b.audit["run_id"])
    check("determinism: same output_hash", a.audit["output_hash"] == b.audit["output_hash"],
          f"{a.audit['output_hash']} != {b.audit['output_hash']}")
    check("determinism: same status trail", a.status_trail == b.status_trail)


def test_config_driven_validation_threshold():
    base = run_pipeline()
    strict = run_pipeline(validation_min_confidence=0.99)
    check("config: default run is clean (no review)", base.review_required is False)
    check("config: raising validation_min_confidence forces review",
          strict.review_required is True)


def test_config_driven_coercion_tokens():
    check("coerce: $ number", coerce("$125,000,000", "number") == 125000000.0)
    check("coerce: comma integer", coerce("1,200", "integer") == 1200)
    check("coerce: default true token", coerce("Yes", "boolean") is True)
    # Custom (config-driven) tokens change behavior:
    check("coerce: custom true token", coerce("ja", "boolean", ["ja"], ["nee"]) is True)
    check("coerce: token not in custom set -> None", coerce("yes", "boolean", ["ja"], ["nee"]) is None)


def test_grounding_gate():
    docs = [{"filename": "a.pdf", "pages": [{"page": 1, "text": "noise"}], "tables": [], "attachments": []}]
    fields = [{
        "name": "insured_name", "value": "Acme", "confidence": 0.99,
        "source_document": "a.pdf", "source_page": 1, "evidence_snippet": "totally unrelated text",
    }]
    results, review = validate_fields(fields, docs, settings())
    check("validation: ungrounded value -> review", review is True and results[0]["status"] == "review",
          str(results))


def test_config_driven_risk_rules():
    from app.agents import RiskReasoningAgent, AgentContext
    from app.llm import MockLLM
    rr = write_json("nevag_risk_test.json", {
        "rules": [{"attribute": "mfa_enabled", "when": {"equals": True}, "flag": "TESTFLAG", "weight": 5}],
        "bands": [{"level": "elevated", "min_score": 5}, {"level": "standard", "min_score": 0}],
    })
    st = {"extracted_fields": [{"name": "mfa_enabled", "value": True, "status": "approved"}]}
    ctx = AgentContext(settings=settings(risk_rules_path=rr), llm=MockLLM())
    RiskReasoningAgent().run(st, ctx)
    check("config: risk rules from config drive flags/level",
          st["risk_assessment"]["level"] == "elevated" and "TESTFLAG" in st["risk_assessment"]["flags"],
          str(st["risk_assessment"]))


def test_autofill_preserves_formula():
    import openpyxl
    res = run_pipeline()
    out = os.path.join(_OUT, res.submission_id + "_filled.xlsx")
    wb = openpyxl.load_workbook(out)
    ws = wb["Rater"]
    check("autofill: input cell written", ws["B3"].value == 125000000)
    check("autofill: formula preserved", ws["B9"].value == "=B3*0.001")


def test_copilot_grounded():
    agent = NevagAgent(settings=settings())
    res = agent.run(NevagInput(query="x", mode=IntakeMode.UPLOAD, files=list(_FILES)))
    ans = agent.ask(res.submission_id, "what is the annual revenue?")
    check("copilot: grounded answer cites value", "125000000" in ans["answer"] and ans["grounded"],
          ans["answer"])
    miss = agent.ask(res.submission_id, "what is missing?")
    check("copilot: missing-question answered", "missing" in miss["answer"].lower())


def test_registry_complete():
    check("registry: 20 agents cataloged", len(AGENT_REGISTRY) == 20, str(len(AGENT_REGISTRY)))
    cat = describe_registry()
    check("registry: every entry has capabilities + version",
          all(c["capabilities"] and c["version"] for c in cat))


def test_heuristic_fallback_flag():
    from app.document_ai import build_ocr_engine, parse_documents
    from app.llm import MockLLM
    s_off = settings(extract_enable_heuristic_fallback=False)
    docs = parse_documents(list(_FILES), s_off, build_ocr_engine(s_off))
    req = ["insured_name", "annual_revenue"]
    off = extract_fields(req, docs, MockLLM(), s_off)
    on = extract_fields(req, docs, MockLLM(), settings())
    check("config: heuristic fallback OFF -> no fields (mock LLM unusable)", off == [])
    check("config: heuristic fallback ON -> fields extracted", len(on) >= 1)


def test_prompt_from_config():
    default = load_extraction_prompt(settings())
    check("config: extraction prompt loaded from config", "extraction engine" in default.lower())
    custom = write_json("nevag_prompt_test.json", {"extraction_system": "CUSTOM-PROMPT-XYZ"})
    over = load_extraction_prompt(settings(prompts_path=custom))
    check("config: prompt override takes effect", over == "CUSTOM-PROMPT-XYZ")


_META_KEYS = {
    "document_id", "file_name", "page_number", "chunk_id",
    "source_type", "submission_id", "created_at", "evidence_snippet",
}


def test_embedding_deterministic():
    emb = MockEmbedding(64)
    a = emb.embed(["Annual Revenue: $125,000,000"])[0]
    b = emb.embed(["Annual Revenue: $125,000,000"])[0]
    check("embeddings: dim from config", len(a) == 64)
    check("embeddings: deterministic (same text -> same vector)", a == b)
    check("embeddings: provider falls back to mock offline",
          build_embedding_provider(settings(embedding_provider="azure_openai")).name == "mock")


def test_memory_vectorstore():
    emb = MockEmbedding(64)
    vs = MemoryVectorStore()
    texts = ["alpha revenue one hundred", "beta employees fifty"]
    metas = [
        {"chunk_id": "d1:1:0", "document_id": "d1", "submission_id": "subA", "file_name": "a.pdf",
         "page_number": 1, "source_type": "pdf", "created_at": "t", "evidence_snippet": "alpha"},
        {"chunk_id": "d2:1:0", "document_id": "d2", "submission_id": "subB", "file_name": "b.pdf",
         "page_number": 1, "source_type": "pdf", "created_at": "t", "evidence_snippet": "beta"},
    ]
    n = vs.upsert_documents(texts, emb.embed(texts), metas)
    check("memory: upsert count", n == 2)
    top = vs.search(emb.embed(["alpha revenue one hundred"])[0], top_k=1)
    check("memory: nearest match returned", top and top[0]["payload"]["document_id"] == "d1", str(top))
    filt = vs.search(emb.embed(["x"])[0], top_k=5, filters={"submission_id": "subB"})
    check("memory: filter by submission", all(h["payload"]["submission_id"] == "subB" for h in filt) and filt)
    removed = vs.delete_by_document_id("d1")
    check("memory: delete_by_document_id", removed == 1 and vs.health_check()["points"] == 1)


def test_qdrant_local_memory():
    try:
        from qdrant_client import QdrantClient
    except Exception as exc:
        check("qdrant: client import (skipped)", True, f"qdrant-client not installed: {exc}")
        return
    client = QdrantClient(":memory:")  # real local Qdrant, no server
    s = settings(vector_store="qdrant", qdrant_url="http://local", qdrant_collection="nevag_test", embedding_dim=64)
    store = QdrantVectorStore(s, dim=64, client=client)
    emb = MockEmbedding(64)
    texts = ["alpha revenue", "beta employees"]
    metas = [
        {"chunk_id": "d1:1:0", "document_id": "d1", "submission_id": "subA", "file_name": "a.pdf",
         "page_number": 1, "source_type": "pdf", "created_at": "t", "evidence_snippet": "alpha"},
        {"chunk_id": "d2:1:0", "document_id": "d2", "submission_id": "subB", "file_name": "b.pdf",
         "page_number": 1, "source_type": "pdf", "created_at": "t", "evidence_snippet": "beta"},
    ]
    store.upsert_documents(texts, emb.embed(texts), metas)
    check("qdrant: dynamic collection created (config name)", client.collection_exists("nevag_test"))
    top = store.search(emb.embed(["alpha revenue"])[0], top_k=1)
    check("qdrant: nearest match", top and top[0]["payload"]["document_id"] == "d1", str(top))
    filt = store.search(emb.embed(["x"])[0], top_k=5, filters={"submission_id": "subB"})
    check("qdrant: filter by submission", bool(filt) and all(h["payload"]["submission_id"] == "subB" for h in filt))
    store.delete_by_document_id("d1")
    hc = store.health_check()
    check("qdrant: health_check enabled + collection listed", hc["enabled"] and "nevag_test" in hc["collections"], str(hc))


def test_factory_and_startup_validation():
    # qdrant requested but no URL -> disabled -> memory fallback
    v_off = validate_vector_store_config(settings(vector_store="qdrant", qdrant_url=None))
    check("startup: qdrant without URL is disabled", v_off["enabled"] is False and "URL" in v_off["reason"])
    store = build_vector_store(settings(vector_store="qdrant", qdrant_url=None), 64)
    check("factory: falls back to memory when qdrant disabled", store.backend == "memory")
    # qdrant URL but no API key -> enabled, unauthenticated (local allowed)
    v_unauth = validate_vector_store_config(settings(vector_store="qdrant", qdrant_url="http://localhost:6333"))
    check("startup: qdrant URL w/o key -> unauthenticated enabled",
          v_unauth["enabled"] is True and "unauthenticated" in v_unauth["reason"])


def test_ingest_and_retrieve():
    s = settings()
    emb = MockEmbedding(64)
    vs = MemoryVectorStore()
    docs = parse_documents(list(_FILES), s, build_ocr_engine(s))
    report = ingest_documents(vs, emb, docs, "subZ", s)
    check("ingest: chunks upserted", report["chunks"] > 0, str(report))
    hits = retrieve(vs, emb, "Annual Revenue", top_k=3, filters={"submission_id": "subZ"})
    check("retrieve: returns hits", len(hits) > 0)
    if hits:
        payload = hits[0]["payload"]
        check("retrieve: full metadata present on chunk", _META_KEYS.issubset(payload.keys()),
              str(sorted(payload.keys())))


def test_service_retrieve():
    agent = NevagAgent(settings=settings())
    res = agent.run(NevagInput(query="x", mode=IntakeMode.UPLOAD, files=list(_FILES)))
    check("service: ingestion report on result", res.retrieval and res.retrieval["chunks"] > 0,
          str(res.retrieval))
    hits = agent.retrieve(res.submission_id, "annual revenue")
    check("service: retrieve scoped to submission",
          bool(hits) and all(h["payload"]["submission_id"] == res.submission_id for h in hits))


def test_bow_embedding_lexical():
    from app.embeddings import MockEmbedding
    from app.vectorstore import _cosine
    emb = MockEmbedding(256)
    q = emb.embed(["annual revenue"])[0]
    near = emb.embed(["annual revenue total gross"])[0]   # shares tokens
    far = emb.embed(["zzz qqq wxyz"])[0]                   # disjoint
    check("bow: lexical overlap -> higher cosine than disjoint",
          _cosine(q, near) > _cosine(q, far))


def test_embedding_cache_reuses():
    from app.embeddings import CachedEmbedding

    class Counting:
        name = "counting"
        dim = 16
        def __init__(self): self.calls = 0
        def is_available(self): return True
        def embed(self, texts):
            self.calls += len(texts)
            return [[0.0] * self.dim for _ in texts]

    inner = Counting()
    cached = CachedEmbedding(inner)
    cached.embed(["a", "b", "a"])   # 'a' duplicated
    cached.embed(["a", "c"])         # 'a' cached, 'c' new
    check("cache: only unique texts embedded", inner.calls == 3, f"calls={inner.calls}")


def test_rag_extraction_and_citations():
    s = settings(rag_extraction_enabled=True, embedding_provider="mock")
    agent = NevagAgent(settings=s)
    res = agent.run(NevagInput(query="New D&O submission", mode=IntakeMode.UPLOAD, files=list(_FILES)))
    approved = [f for f in res.fields if f.status.value == "approved"]
    check("rag: extraction still finds fields", len(res.fields) >= 5, f"{len(res.fields)} fields")
    # at least one field came through the RAG path with citations
    rag_fields = [f for f in res.fields if (f.notes or "") or True]  # inspect raw via audit
    cited = [d for d in res.audit["decisions"] if d.get("retrieval_method") == "rag" and d.get("citations")]
    check("rag: fields carry retrieval citations", len(cited) >= 1, f"cited={len(cited)}")
    check("rag: trail records [rag] mode", any("[rag]" in l for l in res.status_trail))


def test_rag_determinism():
    s = settings(rag_extraction_enabled=True, embedding_provider="mock")
    a = NevagAgent(settings=s).run(NevagInput(query="x", mode=IntakeMode.UPLOAD, files=list(_FILES)))
    b = NevagAgent(settings=s).run(NevagInput(query="x", mode=IntakeMode.UPLOAD, files=list(_FILES)))
    check("rag: deterministic output_hash with RAG on",
          a.audit["output_hash"] == b.audit["output_hash"],
          f"{a.audit['output_hash']} != {b.audit['output_hash']}")


def test_rag_disabled_by_default():
    s = settings()  # rag_extraction_enabled defaults False
    check("rag: disabled by default", s.rag_extraction_enabled is False)


def test_bm25_keyword_index():
    from app.retrieval import KeywordIndex
    ki = KeywordIndex(k1=1.5, b=0.75)
    ki.add("subA", [
        {"text": "Annual Revenue is one hundred million", "payload": {"chunk_id": "c1"}},
        {"text": "Employee headcount details", "payload": {"chunk_id": "c2"}},
    ])
    hits = ki.search("subA", "annual revenue", top_k=2)
    check("bm25: keyword hit ranks first", hits and hits[0]["payload"]["chunk_id"] == "c1", str(hits))
    check("bm25: isolation (unknown submission -> empty)", ki.search("subZ", "revenue", 5) == [])


def test_rrf_fusion():
    from app.retrieval import reciprocal_rank_fusion
    dense = [{"id": "a", "payload": {"chunk_id": "a"}}, {"id": "b", "payload": {"chunk_id": "b"}}]
    kw = [{"id": "b", "payload": {"chunk_id": "b"}}, {"id": "c", "payload": {"chunk_id": "c"}}]
    fused = reciprocal_rank_fusion([dense, kw], k=60)
    check("rrf: doc in both lists ranks first", fused[0]["payload"]["chunk_id"] == "b", str(fused))
    check("rrf: union of all docs", {h["payload"]["chunk_id"] for h in fused} == {"a", "b", "c"})


def test_mmr_diversifies():
    from app.retrieval import mmr
    from app.embeddings import MockEmbedding
    emb = MockEmbedding(128)
    # two near-duplicates + one distinct; MMR with low lambda should include the distinct one
    cands = [
        {"id": "d1", "payload": {"text": "annual revenue one hundred million"}},
        {"id": "d2", "payload": {"text": "annual revenue one hundred million dollars"}},
        {"id": "d3", "payload": {"text": "multi factor authentication enabled"}},
    ]
    qv = emb.embed(["annual revenue"])[0]
    picked = mmr(qv, cands, emb, lam=0.3, top_k=2)
    ids = {p["id"] for p in picked}
    check("mmr: diversifies (includes a distinct source)", "d3" in ids, str(ids))


def test_reranker_heuristic():
    from app.rerank import build_reranker
    r = build_reranker(settings(rerank_provider="heuristic"))
    hits = [
        {"id": "x", "score": 0.4, "payload": {"text": "totally unrelated content here"}},
        {"id": "y", "score": 0.3, "payload": {"text": "annual revenue figures reported"}},
    ]
    out = r.rerank("annual revenue", hits)
    check("rerank: lexical overlap promotes the relevant chunk", out[0]["id"] == "y", str(out))


def test_hybrid_search_pipeline():
    from app.retrieval import Retriever
    from app.document_ai import build_ocr_engine, parse_documents
    s = settings(retrieval_hybrid=True, mmr_enabled=True, rerank_provider="heuristic")
    agent = NevagAgent(settings=s)
    docs = parse_documents(list(_FILES), s, build_ocr_engine(s))
    agent.retriever.ingest(docs, "subH")
    hits = agent.retriever.search("annual revenue", top_k=3, filters={"submission_id": "subH"})
    check("hybrid: returns scoped results", bool(hits) and all(h["payload"]["submission_id"] == "subH" for h in hits))
    check("hybrid: reranked entries carry rerank_score", any("rerank_score" in h for h in hits))


def test_magnitude_coercion():
    from app.extraction import coerce
    check("magnitude: $125M -> 125,000,000", coerce("$125M", "number") == 125000000.0)
    check("magnitude: 1.2bn -> 1,200,000,000", coerce("1.2bn", "number") == 1200000000.0)
    check("magnitude: 100k integer -> 100,000", coerce("100k", "integer") == 100000)
    check("magnitude: plain comma number unchanged", coerce("1,200", "integer") == 1200)
    check("magnitude: full number no false multiply", coerce("$125,000,000", "number") == 125000000.0)


def test_table_extraction():
    from app.llm import MockLLM
    docs = [{
        "filename": "t.xlsx", "sha1": "dt", "kind": "excel",
        "pages": [{"page": 1, "text": ""}],
        "tables": [{"page": 1, "rows": [["Annual Revenue", "$5M"], ["Employee Count", "1,200"]]}],
        "attachments": [],
    }]
    fields = extract_fields(["annual_revenue", "employee_count"], docs, MockLLM(), settings())
    by = {f["name"]: f["value"] for f in fields}
    check("table: value from table row extracted (with magnitude)", by.get("annual_revenue") == 5000000.0, str(by))
    check("table: second table row extracted", by.get("employee_count") == 1200)


def test_shared_candidate_map_text_and_table():
    from app.extraction import build_candidate_map, load_catalog, spec_for
    s = settings()
    docs = [{
        "filename": "mix.docx", "sha1": "dm", "kind": "word",
        "pages": [{"page": 1, "text": "Annual Revenue: 125000000"}],
        "tables": [{"page": 1, "rows": [["Annual Revenue", "$125,000,000"]]}],
        "attachments": [],
    }]
    specs = [spec_for("annual_revenue", load_catalog(s))]
    cmap = build_candidate_map(specs, docs, s)
    origins = {c["origin"] for c in cmap["annual_revenue"]}
    check("candidate map: gathers both text and table candidates", {"text", "table"} <= origins, str(origins))


def test_magnitude_grounding():
    docs = [{"filename": "a.pdf", "sha1": "dg", "kind": "digital_pdf",
             "pages": [{"page": 1, "text": "Annual Revenue: $5M"}], "tables": [], "attachments": []}]
    field = {"name": "annual_revenue", "value": 5000000, "confidence": 0.99,
             "source_document": "a.pdf", "source_page": 1, "evidence_snippet": "Annual Revenue: $5M"}
    res, _ = validate_fields([field], docs, settings())
    check("grounding: magnitude value grounds against $5M snippet (approved)",
          res[0]["status"] == "approved", str(res))


def test_source_priority_resolution():
    from app.llm import MockLLM
    docs = [
        {"filename": "financials_2025.pdf", "sha1": "f", "kind": "digital_pdf",
         "pages": [{"page": 1, "text": "Annual Revenue: 200000000"}], "tables": [], "attachments": []},
        {"filename": "broker_email.eml", "sha1": "e", "kind": "email",
         "pages": [{"page": 1, "text": "Annual Revenue: 100000000"}], "tables": [], "attachments": []},
    ]
    # Priority ON: the financials source wins and the conflict is resolved.
    s_on = settings()
    fields_on = extract_fields(["annual_revenue"], docs, MockLLM(), s_on)
    check("source-priority: authoritative (financials) value chosen",
          fields_on and fields_on[0]["value"] == 200000000.0, str(fields_on))
    res_on, _ = validate_fields(fields_on, docs, s_on)
    check("source-priority: conflict resolved -> approved", res_on[0]["status"] == "approved", str(res_on))
    # Priority OFF: the disagreement is flagged for review.
    s_off = settings(contradiction_use_source_priority=False)
    fields_off = extract_fields(["annual_revenue"], docs, MockLLM(), s_off)
    res_off, _ = validate_fields(fields_off, docs, s_off)
    check("source-priority OFF: contradiction flagged (review)", res_off[0]["status"] == "review", str(res_off))


def test_metrics_surfaced():
    res = run_pipeline()
    m = res.metrics or {}
    check("metrics: total_ms present", "total_ms" in m and isinstance(m["total_ms"], (int, float)))
    check("metrics: counts present", m.get("counts", {}).get("documents", 0) > 0, str(m.get("counts")))
    check("metrics: per-node timings recorded", bool(m.get("node_ms")))


def test_component_health_introspection():
    agent = NevagAgent(settings=settings())
    check("health: llm exposes is_available", hasattr(agent.llm, "is_available"))
    check("health: ocr exposes availability", hasattr(agent._ocr, "available"))
    check("health: embedder is_available callable", agent.embedder.is_available() in (True, False))


def test_parallel_doc_ai_deterministic():
    from app.document_ai import build_ocr_engine, parse_documents
    s1 = settings(doc_ai_max_workers=1)
    s4 = settings(doc_ai_max_workers=4)
    ocr = build_ocr_engine(s1)
    key = lambda docs: [(d["filename"], d["kind"], d["parse_status"], d.get("text", "")) for d in docs]
    r1 = parse_documents(list(_FILES), s1, ocr)
    r4 = parse_documents(list(_FILES), s4, ocr)
    check("parallel doc-ai: same result as sequential (deterministic)", key(r1) == key(r4))


def test_slim_persisted_state():
    agent = NevagAgent(settings=settings())
    res = agent.run(NevagInput(query="x", mode=IntakeMode.UPLOAD, files=list(_FILES)))
    stored = agent.store.load(res.submission_id)
    check("slim state: raw files dropped", "files" not in stored)
    docs = stored.get("parsed_documents", [])
    check("slim state: doc full text dropped, page_count kept",
          bool(docs) and "pages" not in docs[0] and "page_count" in docs[0], str(docs[:1]))


def test_mmr_reuses_stored_vectors():
    from app.retrieval import mmr

    class _NoEmbed:
        name, dim = "noembed", 4
        def is_available(self): return True
        def embed(self, texts): raise AssertionError("MMR should reuse stored vectors, not embed")

    cands = [
        {"id": "a", "vector": [1.0, 0, 0, 0], "payload": {"text": "x"}},
        {"id": "b", "vector": [0, 1.0, 0, 0], "payload": {"text": "y"}},
    ]
    picked = mmr([1.0, 0, 0, 0], cands, _NoEmbed(), 0.5, 2)
    check("mmr: reuses stored vectors (no re-embed)", len(picked) == 2)


def test_sqlite_embedding_cache_persists():
    from app.embeddings import CachedEmbedding, SqliteVectorCache

    class Counting:
        name, dim = "c", 4
        def __init__(self): self.calls = 0
        def is_available(self): return True
        def embed(self, texts): self.calls += len(texts); return [[0.0] * 4 for _ in texts]

    path = os.path.join(tempfile.gettempdir(), "nevag_emb_cache_test.db")
    if os.path.exists(path):
        os.remove(path)
    inner1 = Counting()
    CachedEmbedding(inner1, SqliteVectorCache(path)).embed(["a", "b"])
    inner2 = Counting()
    CachedEmbedding(inner2, SqliteVectorCache(path)).embed(["a", "c"])  # 'a' from disk
    check("emb cache: first instance embeds both", inner1.calls == 2)
    check("emb cache: second instance reuses persisted 'a'", inner2.calls == 1, f"calls={inner2.calls}")


def test_per_attribute_rag_prompts():
    agent = NevagAgent(settings=settings(
        rag_extraction_enabled=True, embedding_provider="mock", rag_per_attribute_prompts=True))
    res = agent.run(NevagInput(query="x", mode=IntakeMode.UPLOAD, files=list(_FILES)))
    rag = [d for d in res.audit["decisions"] if d.get("retrieval_method") == "rag"]
    check("per-attribute RAG: fields extracted with rag method", len(res.fields) >= 4 and len(rag) >= 1,
          f"fields={len(res.fields)} rag={len(rag)}")


def test_duplicate_attribute_reconciliation():
    from app.extraction import reconcile_fields
    # Two materially-different revenue values for the SAME attribute (FRoSTA case).
    dup = [
        {"name": "annual_revenue", "value": 639480000, "confidence": 1.0, "status": "approved",
         "source_document": "Halbjahresfinanzbericht.pdf", "source_type": "digital_pdf"},
        {"name": "annual_revenue", "value": 315941000, "confidence": 1.0, "status": "approved",
         "source_document": "Halbjahresfinanzbericht.pdf", "source_type": "digital_pdf"},
        {"name": "insured_name", "value": "FRoSTA AG", "confidence": 1.0, "status": "approved",
         "source_document": "email.eml", "source_type": "email"},
    ]
    out = reconcile_fields(dup, settings())
    names = [f["name"] for f in out]
    check("reconcile: collapses to one field per attribute", names.count("annual_revenue") == 1, str(names))
    rev = next(f for f in out if f["name"] == "annual_revenue")
    check("reconcile: materially-different values flagged conflict",
          rev.get("reconciliation", {}).get("conflict") is True, str(rev.get("reconciliation")))
    check("reconcile: conflict records all candidates",
          len(rev["reconciliation"]["candidates"]) == 2)
    # Agreeing duplicates -> merged, no conflict.
    same = [
        {"name": "annual_revenue", "value": 125000000, "confidence": 0.9, "status": "approved",
         "source_document": "a.pdf", "source_type": "digital_pdf"},
        {"name": "annual_revenue", "value": 125000000, "confidence": 0.8, "status": "approved",
         "source_document": "b.docx", "source_type": "word"},
    ]
    out2 = reconcile_fields(same, settings())
    check("reconcile: agreeing duplicates merge without conflict",
          len(out2) == 1 and out2[0]["reconciliation"]["conflict"] is False)


def test_reconciliation_conflict_routes_review():
    docs = [{"filename": "a.pdf", "sha1": "d", "kind": "digital_pdf",
             "pages": [{"page": 1, "text": "Annual Revenue: 639480000 and also 315941000"}],
             "tables": [], "attachments": []}]
    field = {"name": "annual_revenue", "value": 639480000, "confidence": 1.0,
             "source_document": "a.pdf", "source_page": 1, "evidence_snippet": "Annual Revenue: 639480000",
             "reconciliation": {"conflict": True, "chosen": 639480000,
                                "candidates": [{"value": 639480000, "source_document": "a.pdf"},
                                               {"value": 315941000, "source_document": "a.pdf"}]}}
    res, _ = validate_fields([field], docs, settings())
    check("reconcile: conflicting attribute routed to review (not approved)",
          res[0]["status"] == "review", str(res))


def test_locale_number_parsing():
    from app.extraction import coerce
    cases = {
        "639.480": 639480.0,        # de thousands (lone '.', 3 trailing)
        "1.234,56": 1234.56,        # de decimal
        "1,234.56": 1234.56,        # en
        "1.234.567": 1234567.0,     # de grouped (repeated '.')
        "$125,000,000": 125000000.0,
        "639 Tsd": 639000.0,        # de scale unit
        "1,2 Mrd": 1200000000.0,    # de decimal + unit
        "1.2bn": 1200000000.0,      # en decimal + unit
    }
    bad = {k: coerce(k, "number") for k, v in cases.items() if coerce(k, "number") != v}
    check("locale numbers: en/de thousands+decimal+scale-units parsed dynamically", not bad, str(bad))


def test_grounding_scale_and_translation():
    from app.validation import _is_grounded
    from app.common import load_value_tokens
    tt, ft = load_value_tokens()
    kw = dict(scale_tolerant=True, min_sig=2, str_overlap=0.34)
    check("grounding: scaled number grounds via significant digits",
          _is_grounded(639480000, "number", "Umsatz: 639.480 T€", tt, ft, **kw))
    check("grounding: true numeric mismatch still fails",
          not _is_grounded(639480000, "number", "Mitarbeiter: 12", tt, ft, **kw))
    check("grounding: translated/paraphrased string grounds via token overlap",
          _is_grounded("Tiefkuehlkost Lebensmittelproduktion", "string",
                       "Branche: Tiefkuehlkost und Lebensmittel", tt, ft, **kw))
    check("grounding: unrelated string still fails",
          not _is_grounded("Acme", "string", "totally unrelated text", tt, ft, **kw))


def test_classification_uses_extracted_industry():
    from app.classification import classify, load_taxonomy
    tax = load_taxonomy(settings())
    r = classify("", "Tiefkühlkost / Lebensmittelproduktion", tax)
    check("classification: uses extracted (multilingual) industry value",
          r["industry"] == "Tiefkühlkost / Lebensmittelproduktion" and r["product"], str(r))


def test_source_priority_ocr_tiebreak():
    # Same priority class (both "application"), clean digital vs OCR'd scan, different
    # values. The clean digital source should win -> conflict resolved -> approved.
    docs = [{"filename": "application.pdf", "sha1": "a", "kind": "digital_pdf",
             "pages": [{"page": 1, "text": "Annual Revenue: 200000000"}], "tables": [], "attachments": []}]
    field = {"name": "annual_revenue", "value": 200000000.0, "confidence": 0.99,
             "source_document": "application.pdf", "source_page": 1,
             "evidence_snippet": "Annual Revenue: 200000000"}
    cmap_mixed = {"annual_revenue": [
        {"value": 200000000.0, "source_document": "application.pdf", "source_type": "digital_pdf"},
        {"value": 199000000.0, "source_document": "scanned_acord.pdf", "source_type": "scanned_pdf"},
    ]}
    res, _ = validate_fields([dict(field)], docs, settings(), candidate_map=cmap_mixed)
    check("source-priority OCR tie-break: clean digital beats noisy scan -> approved",
          res[0]["status"] == "approved", str(res))
    # Control: both degraded (scanned) -> no clean source -> unresolved -> review.
    cmap_both = {"annual_revenue": [
        {"value": 200000000.0, "source_document": "scan_a.pdf", "source_type": "scanned_pdf"},
        {"value": 199000000.0, "source_document": "scan_b.pdf", "source_type": "scanned_pdf"},
    ]}
    res2, _ = validate_fields([dict(field)], docs, settings(), candidate_map=cmap_both)
    check("source-priority: two degraded sources disagree -> review", res2[0]["status"] == "review", str(res2))


def test_dropdown_value_mapping():
    import openpyxl
    from openpyxl.worksheet.datavalidation import DataValidation
    from app.excel_autofill import fill_rater, load_value_mappings

    tpath = os.path.join(tempfile.gettempdir(), "nevag_dd_template.xlsx")
    opath = os.path.join(tempfile.gettempdir(), "nevag_dd_filled.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rater"
    dv = DataValidation(type="list", formula1='"Yes,No"')
    ws.add_data_validation(dv)
    dv.add(ws["B6"])
    wb.save(tpath)
    plan = {"sheet": "Rater", "cells": {"mfa_enabled": {"cell": "B6", "value": True}}}
    fill_rater(tpath, plan, opath, load_value_mappings(settings()))
    got = openpyxl.load_workbook(opath)["Rater"]["B6"].value
    check("dropdown mapping: boolean True -> 'Yes' (carrier vocabulary)", got == "Yes", f"B6={got!r}")


def test_ocr_confidence_threaded():
    from app.document_ai import NullOCR
    from app.retrieval import chunk_documents
    t, c = NullOCR().image_to_text(b"")
    check("ocr engine returns (text, confidence) tuple", t == "" and c == 0.0)
    docs = [{"filename": "s.pdf", "sha1": "d", "kind": "scanned_pdf",
             "pages": [{"page": 1, "text": "Named Insured: X", "ocr": True, "ocr_confidence": 0.87}],
             "tables": [], "attachments": []}]
    chunks = chunk_documents(docs, "subO", settings())
    check("ocr confidence propagated to chunk metadata",
          any(ch["metadata"].get("ocr") and ch["metadata"].get("ocr_confidence") == 0.87 for ch in chunks))


def test_sharepoint_intake_config():
    from app.sharepoint import SharePointClient
    check("sharepoint: unconfigured -> is_configured False", SharePointClient(settings()).is_configured() is False)
    s = settings(sharepoint_site_url="https://x.sharepoint.com", sharepoint_site_name="Site",
                 sharepoint_library_name="Documents", sharepoint_client_id="id",
                 sharepoint_client_secret="sec", sharepoint_tenant_id="t",
                 sharepoint_token_url="https://acs/tokens/OAuth/2")
    c = SharePointClient(s)
    check("sharepoint: configured -> is_configured True", c.is_configured() is True)
    check("sharepoint: server-relative folder built correctly",
          c._folder_server_relative() == "/sites/Site/Documents", c._folder_server_relative())
    check("sharepoint: site host parsed", c._site_host() == "x.sharepoint.com")


def test_perplexity_research_connector():
    from app.research import build_research_sources
    check("research: no key -> no sources", build_research_sources(settings()) == [])
    s = settings(perplexity_api_key="dummy", perplexity_model="sonar-pro")
    srcs = build_research_sources(s)
    check("research: key present -> perplexity source available",
          len(srcs) == 1 and srcs[0].name == "perplexity" and srcs[0].available)
    # No insured in context -> returns None WITHOUT any network call (never invents).
    check("research: no insured -> None (no call, no fabrication)",
          srcs[0].lookup("annual_revenue", {}) is None)


def test_local_stack_defaults():
    from app.llm import OllamaLLM, MockLLM, build_llm
    from app.embeddings import build_embedding_provider
    d = Settings()  # bare defaults
    check("local: default llm_provider is ollama", d.llm_provider == "ollama", d.llm_provider)
    check("local: default embedding_provider is ollama", d.embedding_provider == "ollama")
    check("local: build_llm(ollama) -> OllamaLLM", isinstance(build_llm(Settings(llm_provider="ollama")), OllamaLLM))
    check("local: build_llm(mock) -> MockLLM", isinstance(build_llm(Settings(llm_provider="mock")), MockLLM))
    # Ollama not running -> embeddings degrade to deterministic mock (offline floor)
    emb = build_embedding_provider(Settings(embedding_provider="ollama"))
    check("local: ollama embeddings unavailable -> mock fallback", emb.name == "mock", emb.name)


def test_sqlite_state_store():
    from app.store import SqliteStateStore, build_state_store
    path = os.path.join(tempfile.gettempdir(), "nevag_state_test.db")
    if os.path.exists(path):
        os.remove(path)
    url = "sqlite:///" + path.replace("\\", "/")
    store = build_state_store(settings(database_url=url))
    check("sqlite: factory selects SqliteStateStore", isinstance(store, SqliteStateStore))
    store.save("subS", {"submission_id": "subS", "x": 1})
    check("sqlite: round-trip persists state", store.load("subS") == {"submission_id": "subS", "x": 1})
    check("sqlite: exists()", store.exists("subS") and not store.exists("nope"))


def test_validation_evidence_supply():
    # A field with NO snippet; with retrieval evidence-supply ON it should ground.
    from app.retrieval import Retriever
    from app.embeddings import MockEmbedding
    from app.vectorstore import MemoryVectorStore
    docs = [{"filename": "a.pdf", "sha1": "doc1", "kind": "digital_pdf",
             "pages": [{"page": 1, "text": "Annual Revenue: 125000000"}], "tables": [], "attachments": []}]
    field = {"name": "annual_revenue", "value": 125000000, "confidence": 0.99,
             "source_document": None, "evidence_snippet": None}

    s_on = settings(validation_use_retrieval=True)
    retr = Retriever(s_on, MockEmbedding(int(s_on.embedding_dim)), MemoryVectorStore())
    retr.ingest(docs, "subV")
    res_on, _ = validate_fields([dict(field)], docs, s_on, retriever=retr, submission_id="subV")
    res_off, _ = validate_fields([dict(field)], docs, settings(), retriever=retr, submission_id="subV")
    check("validation evidence-supply: ON grounds the value (approved)", res_on[0]["status"] == "approved", str(res_on))
    check("validation evidence-supply: OFF leaves it unevidenced (review)", res_off[0]["status"] == "review", str(res_off))


def main():
    tests = [
        test_determinism,
        test_config_driven_validation_threshold,
        test_config_driven_coercion_tokens,
        test_grounding_gate,
        test_config_driven_risk_rules,
        test_autofill_preserves_formula,
        test_copilot_grounded,
        test_registry_complete,
        test_heuristic_fallback_flag,
        test_prompt_from_config,
        test_embedding_deterministic,
        test_memory_vectorstore,
        test_qdrant_local_memory,
        test_factory_and_startup_validation,
        test_ingest_and_retrieve,
        test_service_retrieve,
        test_bow_embedding_lexical,
        test_embedding_cache_reuses,
        test_rag_extraction_and_citations,
        test_rag_determinism,
        test_rag_disabled_by_default,
        test_bm25_keyword_index,
        test_rrf_fusion,
        test_mmr_diversifies,
        test_reranker_heuristic,
        test_hybrid_search_pipeline,
        test_validation_evidence_supply,
        test_local_stack_defaults,
        test_sqlite_state_store,
        test_magnitude_coercion,
        test_table_extraction,
        test_shared_candidate_map_text_and_table,
        test_magnitude_grounding,
        test_source_priority_resolution,
        test_metrics_surfaced,
        test_component_health_introspection,
        test_parallel_doc_ai_deterministic,
        test_slim_persisted_state,
        test_mmr_reuses_stored_vectors,
        test_sqlite_embedding_cache_persists,
        test_per_attribute_rag_prompts,
        test_source_priority_ocr_tiebreak,
        test_dropdown_value_mapping,
        test_ocr_confidence_threaded,
        test_perplexity_research_connector,
        test_sharepoint_intake_config,
        test_locale_number_parsing,
        test_grounding_scale_and_translation,
        test_classification_uses_extracted_industry,
        test_duplicate_attribute_reconciliation,
        test_reconciliation_conflict_routes_review,
    ]
    for t in tests:
        try:
            t()
        except Exception as exc:  # a crash is a failed check, not an aborted run
            _RESULTS.append(False)
            print("FAIL -", t.__name__, f"  [EXCEPTION: {type(exc).__name__}: {exc}]")

    passed = sum(1 for r in _RESULTS if r)
    total = len(_RESULTS)
    print(f"\n{passed}/{total} checks passed")
    sys.exit(0 if passed == total else 1)


if __name__ == "__main__":
    main()
